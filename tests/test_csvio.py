from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

from src.csvio import clean_cell, read_csv, write_csv, write_csv_and_xlsx, write_xlsx


class CleanCellTest(unittest.TestCase):
    def test_line_breaks_and_runs_of_spaces_collapse(self) -> None:
        self.assertEqual("a b c", clean_cell("a\r\nb\n  c"))
        self.assertEqual("a b", clean_cell("a b"))
        self.assertEqual("", clean_cell(None))
        self.assertEqual("7", clean_cell(7))

    def test_typographic_characters_are_preserved(self) -> None:
        # The fix for Excel is the BOM, not stripping the document's punctuation.
        self.assertEqual("we’re “neutral” — for now", clean_cell("we’re “neutral” — for now"))


class WriteCsvTest(unittest.TestCase):
    def test_writes_bom_single_line_cells_and_reads_back(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sub" / "out.csv"
            write_csv(path, ("Firm", "Full Commentary"), [{"Firm": "A — B", "Full Commentary": "line1\nline2", "extra": "x"}])
            raw = path.read_bytes()
            self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "BOM required for Excel on Windows")
            self.assertNotIn(b"line1\nline2", raw)
            columns, rows = read_csv(path)
            self.assertEqual(["Firm", "Full Commentary"], columns)
            self.assertEqual([{"Firm": "A — B", "Full Commentary": "line1 line2"}], rows)
            # Plain csv with utf-8-sig sees the same thing (no stray BOM in the header).
            with path.open(newline="", encoding="utf-8-sig") as handle:
                self.assertEqual(["Firm", "Full Commentary"], next(csv.reader(handle)))

    def test_read_csv_accepts_bomless_input(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "legacy.csv"
            path.write_text("Firm,View\nA,O\n", encoding="utf-8")
            self.assertEqual((["Firm", "View"], [{"Firm": "A", "View": "O"}]), read_csv(path))


class WriteXlsxTest(unittest.TestCase):
    def test_xlsx_twin_has_same_rows(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            csv_path, xlsx_path = write_csv_and_xlsx(
                Path(temp_dir) / "out.csv",
                ("Firm", "View"),
                [{"Firm": "A", "View": "O"}, {"Firm": "B\nC", "View": "U"}],
            )
            self.assertEqual(".xlsx", xlsx_path.suffix)
            sheet = load_workbook(xlsx_path).active
            values = [[c.value for c in row] for row in sheet.iter_rows()]
            self.assertEqual([["Firm", "View"], ["A", "O"], ["B C", "U"]], values)
            self.assertEqual("A2", sheet.freeze_panes)

    def test_write_xlsx_alone(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = write_xlsx(Path(temp_dir) / "x.xlsx", ("A",), [{"A": "1"}])
            self.assertTrue(path.is_file())


if __name__ == "__main__":
    unittest.main()
